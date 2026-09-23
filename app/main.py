import sys
import os
import csv
import shutil
import sqlite3
import zipfile
import logging
import json
from copy import copy
from datetime import datetime, date
from pathlib import Path

from PySide6.QtCore import Qt, QRectF
from PySide6.QtGui import QAction, QPixmap, QPainter, QPen, QBrush, QFont
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QFormLayout, QLineEdit, QTextEdit, QPushButton, QLabel,
    QComboBox, QTableWidget, QTableWidgetItem, QMessageBox,
    QFileDialog, QDateEdit, QSpinBox, QDoubleSpinBox, QGroupBox,
    QSplitter, QListWidget, QStackedWidget, QDialog, QDialogButtonBox,
    QHeaderView, QAbstractItemView, QCheckBox, QScrollArea
)

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
)
from reportlab.lib.styles import getSampleStyleSheet
from openpyxl import Workbook
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


APP_NAME = "HSE Management System"
APP_VERSION = "1.2.0"

APP_DIR = Path(os.environ.get("LOCALAPPDATA", Path.home())) / APP_NAME
DB_DIR = APP_DIR / "database"
ATTACH_DIR = APP_DIR / "attachments"
REPORT_DIR = APP_DIR / "reports"
BACKUP_DIR = APP_DIR / "backups"
LOG_DIR = APP_DIR / "logs"

for p in [APP_DIR, DB_DIR, ATTACH_DIR, REPORT_DIR, BACKUP_DIR, LOG_DIR]:
    p.mkdir(parents=True, exist_ok=True)

DB_FILE = DB_DIR / "hse.db"

logging.basicConfig(
    filename=LOG_DIR / "application.log",
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s"
)


# ============================================================
# DATABASE
# ============================================================

class Database:

    def __init__(self):
        self.conn = sqlite3.connect(DB_FILE)
        self.conn.row_factory = sqlite3.Row
        self.conn.execute("PRAGMA foreign_keys = ON")
        self.create_tables()

    def create_tables(self):

        self.conn.executescript("""
        CREATE TABLE IF NOT EXISTS settings (
            key TEXT PRIMARY KEY,
            value TEXT
        );

        CREATE TABLE IF NOT EXISTS employees (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_id TEXT,
            name TEXT NOT NULL,
            designation TEXT,
            company TEXT,
            department TEXT,
            contact TEXT,
            active INTEGER DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS companies (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            company_type TEXT,
            contact_person TEXT,
            contact_number TEXT,
            status TEXT DEFAULT 'Active'
        );

        CREATE TABLE IF NOT EXISTS projects (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project TEXT NOT NULL,
            area TEXT,
            location TEXT,
            description TEXT,
            active INTEGER DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS observations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            number TEXT UNIQUE,
            obs_date TEXT,
            obs_time TEXT,
            project TEXT,
            company TEXT,
            location TEXT,
            area TEXT,
            responsible TEXT,
            designation TEXT,
            employee_id TEXT,
            observer TEXT,
            observer_designation TEXT,
            observer_id TEXT,
            observer_company TEXT,
            obs_type TEXT,
            category TEXT,
            subcategory TEXT,
            observation TEXT,
            immediate_action TEXT,
            corrective_action TEXT,
            preventive_action TEXT,
            priority TEXT,
            target_date TEXT,
            status TEXT DEFAULT 'Open',
            closeout_date TEXT,
            closed_by TEXT,
            verified_by TEXT,
            verification_date TEXT,
            closeout_comments TEXT,
            verification TEXT,
            created_at TEXT
        );

        CREATE TABLE IF NOT EXISTS observation_attachments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            observation_id INTEGER,
            file_path TEXT,
            attachment_type TEXT,
            FOREIGN KEY(observation_id)
                REFERENCES observations(id)
                ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS incident_attachments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            incident_id INTEGER NOT NULL,
            file_path TEXT,
            attachment_type TEXT DEFAULT 'Evidence',
            FOREIGN KEY(incident_id)
                REFERENCES incidents(id)
                ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS audit_attachments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            audit_id INTEGER NOT NULL,
            file_path TEXT,
            attachment_type TEXT DEFAULT 'Evidence',
            FOREIGN KEY(audit_id)
                REFERENCES audits(id)
                ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS capa_attachments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            capa_id INTEGER NOT NULL,
            file_path TEXT,
            attachment_type TEXT DEFAULT 'Evidence',
            FOREIGN KEY(capa_id)
                REFERENCES capa(id)
                ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS incidents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            number TEXT UNIQUE,
            incident_date TEXT,
            incident_time TEXT,
            location TEXT,
            project TEXT,
            company TEXT,
            department TEXT,
            activity TEXT,
            incident_type TEXT,
            person_involved TEXT,
            employee_id TEXT,
            designation TEXT,
            supervisor TEXT,
            witnesses TEXT,
            description TEXT,
            immediate_action TEXT,
            consequences TEXT,
            potential_consequences TEXT,
            equipment TEXT,
            investigation_method TEXT,
            why1 TEXT,
            why2 TEXT,
            why3 TEXT,
            why4 TEXT,
            why5 TEXT,
            direct_cause TEXT,
            contributing_factors TEXT,
            root_cause TEXT,
            corrective_action TEXT,
            preventive_action TEXT,
            status TEXT DEFAULT 'Open',
            created_at TEXT
        );

        CREATE TABLE IF NOT EXISTS audits (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            number TEXT UNIQUE,
            audit_date TEXT,
            audit_type TEXT,
            standard TEXT,
            project TEXT,
            location TEXT,
            department TEXT,
            auditor TEXT,
            lead_auditor TEXT,
            auditee TEXT,
            scope TEXT,
            objective TEXT,
            criteria TEXT,
            start_time TEXT,
            end_time TEXT,
            status TEXT DEFAULT 'Open',
            created_at TEXT
        );

        CREATE TABLE IF NOT EXISTS audit_findings (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            audit_id INTEGER,
            clause TEXT,
            sub_clause TEXT,
            requirement TEXT,
            finding_type TEXT,
            observation TEXT,
            evidence TEXT,
            risk_impact TEXT,
            corrective_action TEXT,
            responsible TEXT,
            target_date TEXT,
            status TEXT DEFAULT 'Open',
            verification TEXT,
            closeout_evidence TEXT,
            FOREIGN KEY(audit_id)
                REFERENCES audits(id)
                ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS capa (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            number TEXT UNIQUE,
            source TEXT,
            reference_number TEXT,
            finding TEXT,
            root_cause TEXT,
            corrective_action TEXT,
            preventive_action TEXT,
            responsible TEXT,
            priority TEXT,
            target_date TEXT,
            status TEXT DEFAULT 'Open',
            verification TEXT,
            closeout_evidence TEXT,
            created_at TEXT
        );

        CREATE TABLE IF NOT EXISTS categories (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            category TEXT,
            subcategory TEXT,
            active INTEGER DEFAULT 1
        );
        """)

        # Backward-compatible migrations for existing installations.
        self.conn.execute("CREATE TABLE IF NOT EXISTS incident_attachments (id INTEGER PRIMARY KEY AUTOINCREMENT, incident_id INTEGER NOT NULL, file_path TEXT, attachment_type TEXT DEFAULT 'Evidence', FOREIGN KEY(incident_id) REFERENCES incidents(id) ON DELETE CASCADE)")
        self.conn.execute("CREATE TABLE IF NOT EXISTS audit_attachments (id INTEGER PRIMARY KEY AUTOINCREMENT, audit_id INTEGER NOT NULL, file_path TEXT, attachment_type TEXT DEFAULT 'Evidence', FOREIGN KEY(audit_id) REFERENCES audits(id) ON DELETE CASCADE)")
        self.conn.execute("CREATE TABLE IF NOT EXISTS capa_attachments (id INTEGER PRIMARY KEY AUTOINCREMENT, capa_id INTEGER NOT NULL, file_path TEXT, attachment_type TEXT DEFAULT 'Evidence', FOREIGN KEY(capa_id) REFERENCES capa(id) ON DELETE CASCADE)")
        self.conn.commit()

    def execute(self, sql, params=()):
        cur = self.conn.cursor()
        cur.execute(sql, params)
        self.conn.commit()
        return cur

    def fetchall(self, sql, params=()):
        return self.conn.execute(sql, params).fetchall()

    def fetchone(self, sql, params=()):
        return self.conn.execute(sql, params).fetchone()

    def setting(self, key, default=""):
        row = self.fetchone(
            "SELECT value FROM settings WHERE key=?",
            (key,)
        )
        return row["value"] if row else default

    def set_setting(self, key, value):
        self.execute(
            """
            INSERT INTO settings(key,value)
            VALUES(?,?)
            ON CONFLICT(key)
            DO UPDATE SET value=excluded.value
            """,
            (key, value)
        )


db = Database()


# ============================================================
# HELPERS
# ============================================================

OBS_TYPES = [
    "Unsafe Act",
    "Unsafe Condition",
    "Good Observation",
    "Good Practice",
    "Positive Observation",
    "Environmental Observation",
    "Near Miss",
    "Other"
]

CATEGORIES = [
    "PPE",
    "Work at Height",
    "Excavation",
    "Electrical Safety",
    "Lifting Operations",
    "Scaffolding",
    "Confined Space",
    "Fire Safety",
    "Housekeeping",
    "Slip Trip Fall",
    "Chemical Safety",
    "Environmental",
    "Traffic Management",
    "Plant and Machinery",
    "Manual Handling",
    "Emergency Preparedness",
    "Heat Stress",
    "Welfare",
    "Dropped Objects",
    "Tools and Equipment",
    "Permit to Work",
    "Barricading",
    "Material Storage",
    "Access and Egress",
    "Gas Testing",
    "Hot Work",
    "Noise",
    "Dust",
    "Waste Management",
    "Environmental Spill",
    "Other"
]

PRIORITIES = ["Low", "Medium", "High", "Critical"]

STATUSES = [
    "Open",
    "In Progress",
    "Pending Verification",
    "Closed",
    "Cancelled"
]

INCIDENT_TYPES = [
    "Fatality",
    "Lost Time Injury",
    "Medical Treatment Case",
    "Restricted Work Case",
    "First Aid Case",
    "Near Miss",
    "Property Damage",
    "Environmental Incident",
    "Fire Incident",
    "Vehicle Incident",
    "Equipment Damage",
    "Chemical Spill",
    "Other"
]

INVESTIGATION_METHODS = [
    "Root Cause Analysis",
    "5 Why Analysis",
    "Fishbone / Ishikawa",
    "ICAM",
    "Barrier Analysis",
    "Bow-Tie Analysis",
    "Fault Tree Analysis",
    "Causal Tree",
    "Other"
]

AUDIT_TYPES = [
    "Internal Audit",
    "External Audit",
    "Client Audit",
    "Certification Audit",
    "Surveillance Audit",
    "Regulatory Audit",
    "Project Audit",
    "Supplier Audit",
    "Other"
]

FINDING_TYPES = [
    "Positive Observation",
    "Good Practice",
    "Conformity",
    "Opportunity for Improvement",
    "Observation",
    "Minor Nonconformity",
    "Major Nonconformity",
    "Environmental Finding",
    "Legal / Compliance Finding",
    "Other"
]

CAPA_SOURCES = [
    "HSE Inspection",
    "Incident",
    "Near Miss",
    "Audit",
    "Client Observation",
    "Regulatory Inspection",
    "Environmental Incident",
    "Management Review",
    "Employee Observation",
    "Other"
]


def next_number(prefix, table):
    year = datetime.now().year
    row = db.fetchone(
        f"""
        SELECT number FROM {table}
        WHERE number LIKE ?
        ORDER BY id DESC LIMIT 1
        """,
        (f"{prefix}-{year}-%",)
    )

    if not row:
        n = 1
    else:
        try:
            n = int(row["number"].split("-")[-1]) + 1
        except Exception:
            n = 1

    return f"{prefix}-{year}-{n:05d}"


def today():
    return datetime.now().strftime("%Y-%m-%d")


def now_time():
    return datetime.now().strftime("%H:%M")


def safe(value):
    return "" if value is None else str(value)


def overdue(target, status):
    if not target or status in ("Closed", "Cancelled"):
        return False
    try:
        return date.fromisoformat(target) < date.today()
    except Exception:
        return False



# ============================================================
# DOCUMENT / ATTACHMENT HELPERS
# ============================================================

ISO_CLAUSES = {
    "ISO 45001": {
        "4": ["4.1 Understanding the organization and its context",
              "4.2 Understanding the needs and expectations of workers and other interested parties",
              "4.3 Determining the scope of the OH&S management system",
              "4.4 OH&S management system"],
        "5": ["5.1 Leadership and commitment", "5.2 OH&S policy",
              "5.3 Organizational roles, responsibilities and authorities",
              "5.4 Consultation and participation of workers"],
        "6": ["6.1 Actions to address risks and opportunities",
              "6.1.2 Hazard identification and assessment of risks and opportunities",
              "6.1.3 Determination of legal requirements and other requirements",
              "6.2 OH&S objectives and planning to achieve them"],
        "7": ["7.1 Resources", "7.2 Competence", "7.3 Awareness",
              "7.4 Communication", "7.5 Documented information"],
        "8": ["8.1 Operational planning and control", "8.1.2 Eliminating hazards and reducing OH&S risks",
              "8.2 Emergency preparedness and response"],
        "9": ["9.1 Monitoring, measurement, analysis and performance evaluation",
              "9.2 Internal audit", "9.3 Management review"],
        "10": ["10.1 General", "10.2 Incident, nonconformity and corrective action",
               "10.3 Continual improvement"],
    },
    "ISO 14001": {
        "4": ["4.1 Understanding the organization and its context",
              "4.2 Understanding the needs and expectations of interested parties",
              "4.3 Determining the scope of the environmental management system",
              "4.4 Environmental management system"],
        "5": ["5.1 Leadership and commitment", "5.2 Environmental policy",
              "5.3 Organizational roles, responsibilities and authorities"],
        "6": ["6.1 Actions to address risks and opportunities",
              "6.1.2 Environmental aspects", "6.1.3 Compliance obligations",
              "6.1.4 Planning action", "6.2 Environmental objectives and planning"],
        "7": ["7.1 Resources", "7.2 Competence", "7.3 Awareness",
              "7.4 Communication", "7.5 Documented information"],
        "8": ["8.1 Operational planning and control", "8.2 Emergency preparedness and response"],
        "9": ["9.1 Monitoring, measurement, analysis and evaluation",
              "9.2 Internal audit", "9.3 Management review"],
        "10": ["10.1 General", "10.2 Nonconformity and corrective action",
               "10.3 Continual improvement"],
    }
}


def company_name():
    return db.setting("company_name", "") or APP_NAME


def document_prefix():
    return db.setting("document_prefix", "HSE") or "HSE"


def document_number(kind):
    prefix = document_prefix()
    return next_number(f"{prefix}-{kind}", {
        "INC": "incidents",
        "OBS": "observations",
        "AUD": "audits",
        "CAPA": "capa"
    }.get(kind, "incidents"))


def attachment_rows(table_name, record_id):
    allowed = {
        "observation_attachments": "observation_id",
        "incident_attachments": "incident_id",
        "audit_attachments": "audit_id",
        "capa_attachments": "capa_id"
    }
    if table_name not in allowed:
        return []
    return db.fetchall(
        f"SELECT * FROM {table_name} WHERE {allowed[table_name]}=? ORDER BY id",
        (record_id,)
    )


def copy_attachments(paths, record_number, table_name, record_id, attachment_type="Evidence"):
    column = {
        "observation_attachments": "observation_id",
        "incident_attachments": "incident_id",
        "audit_attachments": "audit_id",
        "capa_attachments": "capa_id"
    }[table_name]
    saved = []
    for path in paths or []:
        try:
            source = Path(path)
            if not source.exists():
                continue
            destination = ATTACH_DIR / f"{record_number}_{source.name}"
            counter = 1
            while destination.exists():
                destination = ATTACH_DIR / f"{record_number}_{counter}_{source.name}"
                counter += 1
            shutil.copy2(source, destination)
            db.execute(
                f"INSERT INTO {table_name} ({column},file_path,attachment_type) VALUES(?,?,?)",
                (record_id, str(destination), attachment_type)
            )
            saved.append(str(destination))
        except Exception:
            logging.exception("Failed to save attachment %s", path)
    return saved


def add_docx_cell_shading(cell, fill="17365D"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_docx_cell_borders(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def style_docx_table(table, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_docx_cell_borders(
                cell,
                top={"val": "single", "sz": 6, "color": "808080"},
                bottom={"val": "single", "sz": 6, "color": "808080"},
                left={"val": "single", "sz": 6, "color": "808080"},
                right={"val": "single", "sz": 6, "color": "808080"},
            )
            if header and row_index == 0:
                add_docx_cell_shading(cell)
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = __import__("docx").shared.RGBColor(255,255,255)


def add_docx_header(document):
    section = document.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo = db.setting("company_logo", "")
    if logo and Path(logo).exists():
        try:
            run = p.add_run()
            run.add_picture(logo, width=Inches(1.15))
        except Exception:
            pass
    run = p.add_run(company_name())
    run.bold = True
    run.font.size = Pt(14)


def add_docx_footer(document):
    footer_text = db.setting("report_footer", "")
    if not footer_text:
        return
    p = document.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(footer_text).font.size = Pt(8)


def add_docx_title(document, title, number=""):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(18)
    if number:
        p2 = document.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(f"Document No.: {number}")
        r2.bold = True
        r2.font.size = Pt(10)


def add_docx_kv_table(document, pairs):
    table = document.add_table(rows=0, cols=2)
    for key, value in pairs:
        cells = table.add_row().cells
        cells[0].text = safe(key)
        cells[1].text = safe(value)
        cells[0].paragraphs[0].runs[0].bold = True
    style_docx_table(table, header=False)
    return table


def add_docx_attachments(document, rows):
    document.add_heading("Evidence / Attachments", level=2)
    if not rows:
        document.add_paragraph("No attachments recorded.")
        return
    table = document.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "No."
    hdr[1].text = "File"
    hdr[2].text = "Type"
    for i, row in enumerate(rows, 1):
        cells = table.add_row().cells
        cells[0].text = str(i)
        cells[1].text = Path(safe(row["file_path"])).name
        cells[2].text = safe(row["attachment_type"])
    style_docx_table(table)


def add_docx_signatures(document):
    document.add_paragraph()
    document.add_heading("Review and Approval", level=2)
    table = document.add_table(rows=2, cols=3)
    labels = ["Prepared / HSE Report Made By", "Reviewed By", "Approved By"]
    for i, label in enumerate(labels):
        table.cell(0, i).text = label
        table.cell(1, i).text = "\n\n____________________________\nSignature: __________________\nDate: ______________________"
    style_docx_table(table, header=True)


def generate_incident_docx(incident_id, path):
    row = db.fetchone("SELECT * FROM incidents WHERE id=?", (incident_id,))
    if not row:
        raise ValueError("Incident record not found.")

    doc = Document()
    add_docx_header(doc)
    add_docx_title(doc, "ACCIDENT / INCIDENT INVESTIGATION REPORT", row["number"])

    add_docx_kv_table(doc, [
        ("Company", company_name()),
        ("Project", row["project"]),
        ("Incident Date", row["incident_date"]),
        ("Incident Time", row["incident_time"]),
        ("Location", row["location"]),
        ("Department", row["department"]),
        ("Activity", row["activity"]),
        ("Incident Type", row["incident_type"]),
        ("Person Involved", row["person_involved"]),
        ("Employee ID", row["employee_id"]),
        ("Designation", row["designation"]),
        ("Supervisor", row["supervisor"]),
        ("Witnesses", row["witnesses"]),
        ("Equipment", row["equipment"]),
        ("Investigation Method", row["investigation_method"]),
    ])

    sections = [
        ("Incident Description", row["description"]),
        ("Immediate Action", row["immediate_action"]),
        ("Actual Consequences", row["consequences"]),
        ("Potential Consequences", row["potential_consequences"]),
        ("Direct Cause", row["direct_cause"]),
        ("Contributing Factors", row["contributing_factors"]),
        ("Root Cause", row["root_cause"]),
        ("Corrective Action", row["corrective_action"]),
        ("Preventive Action", row["preventive_action"]),
    ]
    for heading, value in sections:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(safe(value) or "N/A")

    doc.add_heading("5 Why Analysis", level=2)
    why_table = doc.add_table(rows=1, cols=2)
    why_table.rows[0].cells[0].text = "Step"
    why_table.rows[0].cells[1].text = "Analysis"
    for i in range(1, 6):
        cells = why_table.add_row().cells
        cells[0].text = f"Why {i}"
        cells[1].text = safe(row[f"why{i}"]) or "N/A"
    style_docx_table(why_table)

    add_docx_attachments(doc, attachment_rows("incident_attachments", incident_id))
    add_docx_signatures(doc)
    add_docx_footer(doc)
    doc.save(path)


def generate_table_docx(table_name, path):
    columns, data = MainWindow.table_data_static(table_name)
    if not columns:
        raise ValueError("There is no data to export.")
    doc = Document()
    add_docx_header(doc)
    add_docx_title(doc, f"{table_name.replace('_', ' ').title()} Report")
    p = doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    table = doc.add_table(rows=1, cols=len(columns))
    for i, col in enumerate(columns):
        table.rows[0].cells[i].text = str(col)
    for row in data[:1000]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = safe(value)
    style_docx_table(table)
    add_docx_footer(doc)
    doc.save(path)


class PieChartWidget(QWidget):
    def __init__(self, values=None, parent=None):
        super().__init__(parent)
        self.values = values or {}
        self.setMinimumHeight(260)

    def set_values(self, values):
        self.values = values or {}
        self.update()

    def paintEvent(self, event):
        painter=QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        rect=self.rect().adjusted(20,20,-20,-20)
        total=sum(max(0,int(v)) for v in self.values.values())
        if total<=0:
            painter.drawText(rect,Qt.AlignmentFlag.AlignCenter,"No observation data")
            painter.end(); return
        center=rect.center()
        radius=min(rect.width(),rect.height())//3
        start=0
        palette=[Qt.GlobalColor.darkBlue,Qt.GlobalColor.darkGreen,Qt.GlobalColor.darkRed,
                 Qt.GlobalColor.darkCyan,Qt.GlobalColor.darkMagenta,Qt.GlobalColor.darkYellow,
                 Qt.GlobalColor.gray]
        for i,(label,value) in enumerate(self.values.items()):
            value=max(0,int(value))
            span=int(round(360*16*value/total))
            painter.setBrush(QBrush(palette[i%len(palette)]))
            painter.setPen(QPen(Qt.GlobalColor.white,1))
            painter.drawPie(QRectF(center.x()-radius,center.y()-radius,2*radius,2*radius),start,span)
            start += span
        y=25
        x=rect.left()
        painter.setFont(QFont("Arial",9))
        for i,(label,value) in enumerate(self.values.items()):
            painter.setBrush(QBrush(palette[i%len(palette)]))
            painter.drawRect(x,y,14,14)
            painter.setPen(QPen(Qt.GlobalColor.black))
            painter.drawText(x+20,y+12,f"{label}: {value}")
            y += 20
        painter.end()


class BarChartWidget(QWidget):
    def __init__(self, values=None, parent=None):
        super().__init__(parent)
        self.values=values or {}
        self.setMinimumHeight(260)

    def set_values(self, values):
        self.values=values or {}
        self.update()

    def paintEvent(self,event):
        painter=QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        rect=self.rect().adjusted(45,20,-20,-45)
        if not self.values:
            painter.drawText(rect,Qt.AlignmentFlag.AlignCenter,"No data")
            painter.end(); return
        maxv=max([int(v) for v in self.values.values()] or [1])
        count=len(self.values)
        barw=max(18,int(rect.width()/max(count,1)*0.65))
        gap=max(8,int(rect.width()/max(count,1)*0.35))
        x=rect.left()
        colors=[Qt.GlobalColor.darkBlue,Qt.GlobalColor.darkGreen,Qt.GlobalColor.darkRed,
                Qt.GlobalColor.darkCyan,Qt.GlobalColor.darkMagenta,Qt.GlobalColor.darkYellow]
        for i,(label,value) in enumerate(self.values.items()):
            value=int(value)
            h=int(rect.height()*value/maxv) if maxv else 0
            y=rect.bottom()-h
            painter.setBrush(QBrush(colors[i%len(colors)]))
            painter.setPen(QPen(Qt.GlobalColor.black))
            painter.drawRect(x,y,barw,h)
            painter.setPen(QPen(Qt.GlobalColor.black))
            painter.drawText(x,y-5,str(value))
            painter.save()
            painter.translate(x+barw/2,rect.bottom()+12)
            painter.rotate(-45)
            painter.drawText(0,0,str(label)[:16])
            painter.restore()
            x += barw+gap
        painter.drawLine(rect.left(),rect.bottom(),rect.right(),rect.bottom())
        painter.end()


# ============================================================
# MAIN WINDOW
# ============================================================

class MainWindow(QMainWindow):

    def __init__(self):
        super().__init__()

        self.setWindowTitle(
            f"{APP_NAME} - {APP_VERSION}"
        )

        self.resize(1400, 850)

        self.setStyleSheet("""
        QMainWindow {
            background: #F4F6F8;
        }

        QLabel {
            color: #202020;
        }

        QPushButton {
            background: #17365D;
            color: white;
            border: none;
            padding: 9px 14px;
            border-radius: 5px;
            min-height: 20px;
        }

        QPushButton:hover {
            background: #245486;
        }

        QLineEdit, QTextEdit, QComboBox, QDateEdit {
            background: white;
            border: 1px solid #B8C2CC;
            border-radius: 4px;
            padding: 7px;
        }

        QTableWidget {
            background: white;
            gridline-color: #D5DADF;
        }

        QHeaderView::section {
            background: #17365D;
            color: white;
            padding: 7px;
            border: none;
        }

        QListWidget {
            background: #102A43;
            color: white;
            border: none;
        }

        QListWidget::item {
            padding: 14px;
        }

        QListWidget::item:selected {
            background: #245486;
        }

        QGroupBox {
            font-weight: bold;
            border: 1px solid #CCD3DA;
            margin-top: 10px;
            padding: 10px;
        }
        """)

        self.build_ui()
        self.dashboard()

    # --------------------------------------------------------

    def build_ui(self):

        central = QWidget()
        self.setCentralWidget(central)

        layout = QHBoxLayout(central)
        layout.setContentsMargins(0, 0, 0, 0)

        self.nav = QListWidget()

        modules = [
            "Dashboard",
            "HSE Inspection Register",
            "Incident Investigation",
            "Audit Register",
            "CAPA Register",
            "Reports",
            "Master Data",
            "Settings"
        ]

        self.nav.addItems(modules)
        self.nav.setFixedWidth(245)
        self.nav.currentRowChanged.connect(
            self.navigation_changed
        )

        layout.addWidget(self.nav)

        self.stack = QStackedWidget()
        layout.addWidget(self.stack, 1)

    # --------------------------------------------------------

    def navigation_changed(self, row):

        if row == 0:
            self.dashboard()
        elif row == 1:
            self.observations()
        elif row == 2:
            self.incidents()
        elif row == 3:
            self.audits()
        elif row == 4:
            self.capa()
        elif row == 5:
            self.reports()
        elif row == 6:
            self.master_data()
        elif row == 7:
            self.settings_page()

    # --------------------------------------------------------

    def page(self, title):

        w = QWidget()
        layout = QVBoxLayout(w)

        title_label = QLabel(title)
        title_label.setStyleSheet("""
            font-size: 24px;
            font-weight: bold;
            color: #17365D;
            padding: 8px;
        """)

        layout.addWidget(title_label)

        self.stack.addWidget(w)
        self.stack.setCurrentWidget(w)

        return w, layout

    # ========================================================
    # DASHBOARD
    # ========================================================


    def dashboard(self):
        w,layout=self.page("HSE Dashboard")
        cards=QHBoxLayout()
        values=[
            ("Observations","SELECT COUNT(*) c FROM observations"),
            ("Open Observations","SELECT COUNT(*) c FROM observations WHERE status NOT IN ('Closed','Cancelled')"),
            ("Overdue","SELECT COUNT(*) c FROM observations WHERE status NOT IN ('Closed','Cancelled') AND target_date < date('now')"),
            ("Incidents","SELECT COUNT(*) c FROM incidents"),
            ("Audits","SELECT COUNT(*) c FROM audits"),
            ("CAPA","SELECT COUNT(*) c FROM capa")
        ]
        for name,sql in values:
            value=db.fetchone(sql)["c"]; box=QGroupBox(name); bl=QVBoxLayout(box)
            label=QLabel(str(value)); label.setStyleSheet("font-size:30px;font-weight:bold;color:#17365D;")
            bl.addWidget(label); cards.addWidget(box)
        layout.addLayout(cards)
        layout.addWidget(QLabel(f"Company: {company_name()}    Project: {db.setting('project_name','')}"))

        charts=QHBoxLayout()
        pie=PieChartWidget()
        obs_types=db.fetchall("SELECT obs_type,COUNT(*) c FROM observations GROUP BY obs_type ORDER BY c DESC")
        pie.set_values({safe(r["obs_type"]) or "Unspecified":r["c"] for r in obs_types})
        pie_box=QGroupBox("Observation Type Distribution"); pl=QVBoxLayout(pie_box); pl.addWidget(pie)
        bar=BarChartWidget()
        stats={"Observations":db.fetchone("SELECT COUNT(*) c FROM observations")["c"],
               "Incidents":db.fetchone("SELECT COUNT(*) c FROM incidents")["c"],
               "Audits":db.fetchone("SELECT COUNT(*) c FROM audits")["c"],
               "CAPA":db.fetchone("SELECT COUNT(*) c FROM capa")["c"]}
        bar.set_values(stats)
        bar_box=QGroupBox("HSE Register Summary"); bl=QVBoxLayout(bar_box); bl.addWidget(bar)
        charts.addWidget(pie_box); charts.addWidget(bar_box); layout.addLayout(charts)

        table=QTableWidget(); table.setColumnCount(6)
        table.setHorizontalHeaderLabels(["Observation","Type","Category","Location","Priority","Status"])
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        rows=db.fetchall("SELECT number,obs_type,category,location,priority,status FROM observations ORDER BY id DESC LIMIT 10")
        table.setRowCount(len(rows))
        for r,row in enumerate(rows):
            vals=[row["number"],row["obs_type"],row["category"],row["location"],row["priority"],row["status"]]
            for c,v in enumerate(vals):
                table.setItem(r,c,QTableWidgetItem(safe(v)))
        layout.addWidget(QLabel("Recent Observations"))
        layout.addWidget(table)


    def observations(self):
        w, layout = self.page("HSE Inspection Register")
        toolbar = QHBoxLayout()
        search = QLineEdit()
        search.setPlaceholderText("Search number, location, responsible person, description...")
        toolbar.addWidget(search)
        add_button = QPushButton("+ New Observation")
        delete_button = QPushButton("Delete Selected")
        export_button = QPushButton("Export CSV")
        excel_button = QPushButton("Export Excel")
        pdf_button = QPushButton("Export PDF")
        word_button = QPushButton("Export Word")
        toolbar.addWidget(add_button)
        toolbar.addWidget(delete_button)
        toolbar.addWidget(export_button)
        toolbar.addWidget(excel_button)
        toolbar.addWidget(pdf_button)
        toolbar.addWidget(word_button)
        layout.addLayout(toolbar)

        table = QTableWidget()
        headers = ["ID","Number","Date","Type","Category","Location","Responsible",
                   "Priority","Target","Status","Overdue","Observation Details"]
        table.setColumnCount(len(headers))
        table.setHorizontalHeaderLabels(headers)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        table.horizontalHeader().setStretchLastSection(True)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setSelectionMode(QAbstractItemView.SingleSelection)
        table.setWordWrap(True)
        layout.addWidget(table)

        def load():
            term = search.text().strip()
            query = "SELECT * FROM observations"
            params = ()
            if term:
                query += """ WHERE number LIKE ? OR location LIKE ? OR responsible LIKE ? OR observation LIKE ?"""
                like = f"%{term}%"
                params = (like, like, like, like)
            query += " ORDER BY id DESC"
            rows = db.fetchall(query, params)
            table.setRowCount(len(rows))
            for r, row in enumerate(rows):
                values = [
                    row["id"], row["number"], row["obs_date"], row["obs_type"],
                    row["category"], row["location"], row["responsible"],
                    row["priority"], row["target_date"], row["status"],
                    "OVERDUE" if overdue(row["target_date"], row["status"]) else "",
                    row["observation"]
                ]
                for c, value in enumerate(values):
                    item = QTableWidgetItem(safe(value))
                    item.setToolTip(safe(value))
                    if c == 10 and value:
                        item.setForeground(Qt.GlobalColor.red)
                    table.setItem(r, c, item)

        def delete_selected():
            row = table.currentRow()
            if row < 0:
                QMessageBox.warning(self, "Delete Observation", "Select an observation first.")
                return
            item = table.item(row, 0)
            if not item:
                return
            obs_id = int(item.text())
            record = db.fetchone("SELECT number FROM observations WHERE id=?", (obs_id,))
            if not record:
                return
            answer = QMessageBox.question(
                self, "Confirm Delete",
                f"Delete observation {record['number']} and its evidence?\n\nThis cannot be undone."
            )
            if answer != QMessageBox.Yes:
                return
            attachments = attachment_rows("observation_attachments", obs_id)
            for a in attachments:
                try:
                    Path(a["file_path"]).unlink(missing_ok=True)
                except Exception:
                    logging.exception("Unable to remove observation attachment")
            db.execute("DELETE FROM observations WHERE id=?", (obs_id,))
            load()

        search.textChanged.connect(load)
        add_button.clicked.connect(lambda: self.observation_form(load))
        delete_button.clicked.connect(delete_selected)
        export_button.clicked.connect(lambda: self.export_csv("observations"))
        excel_button.clicked.connect(lambda: self.export_excel("observations"))
        pdf_button.clicked.connect(lambda: self.export_pdf("observations"))
        word_button.clicked.connect(lambda: self.export_docx("observations"))
        load()



    def observation_form(self, refresh):
        dialog = QDialog(self)
        dialog.setWindowTitle("New HSE Observation / Inspection")
        dialog.resize(980, 820)
        dialog.setMinimumSize(900, 700)
        outer = QVBoxLayout(dialog)

        header = QLabel("HSE Inspection Observation")
        header.setStyleSheet("font-size:20px;font-weight:bold;color:#17365D;padding:6px;")
        outer.addWidget(header)

        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        content = QWidget()
        form = QFormLayout(content)
        form.setFieldGrowthPolicy(QFormLayout.FieldGrowthPolicy.ExpandingFieldsGrow)
        edits = {}

        def edit(name):
            e = QLineEdit()
            edits[name] = e
            form.addRow(name.replace("_", " ").title() + ":", e)
            return e

        for name in ["project", "company", "location", "area", "responsible", "designation", "employee_id",
                     "observer", "observer_designation", "observer_id", "observer_company"]:
            edit(name)

        obs_type = QComboBox(); obs_type.addItems(OBS_TYPES)
        form.addRow("Observation Type:", obs_type)
        category = QComboBox(); category.addItems(CATEGORIES)
        form.addRow("Category:", category)
        edit("subcategory")

        observation = QTextEdit(); observation.setMinimumHeight(100)
        form.addRow("Observation Description:", observation)
        immediate = QTextEdit(); immediate.setMinimumHeight(80)
        form.addRow("Immediate Action:", immediate)
        corrective = QTextEdit(); corrective.setMinimumHeight(80)
        form.addRow("Corrective Action:", corrective)
        preventive = QTextEdit(); preventive.setMinimumHeight(80)
        form.addRow("Preventive Action:", preventive)

        priority = QComboBox(); priority.addItems(PRIORITIES)
        form.addRow("Priority:", priority)
        target = QDateEdit(); target.setCalendarPopup(True); target.setDate(datetime.now().date())
        form.addRow("Target Completion Date:", target)
        status = QComboBox(); status.addItems(STATUSES)
        form.addRow("Status:", status)

        attachment_paths = []
        attachment_label = QLabel("No evidence files selected.")
        attachment_label.setWordWrap(True)
        attach_button = QPushButton("Attach Evidence Files")
        attach_button.setMinimumHeight(36)

        def choose_files():
            paths, _ = QFileDialog.getOpenFileNames(
                dialog, "Select Evidence Files", "",
                "Evidence Files (*.png *.jpg *.jpeg *.bmp *.pdf *.doc *.docx *.xls *.xlsx *.txt *.mp4 *.avi *.mov);;All Files (*)"
            )
            if paths:
                attachment_paths.clear()
                attachment_paths.extend(paths)
                attachment_label.setText("\n".join(Path(p).name for p in paths))

        attach_button.clicked.connect(choose_files)
        form.addRow("Evidence / Attachments:", attach_button)
        form.addRow("Selected Files:", attachment_label)

        scroll_area.setWidget(content)
        outer.addWidget(scroll_area, 1)

        buttons = QDialogButtonBox()
        save_button = buttons.addButton("Save Observation", QDialogButtonBox.ButtonRole.AcceptRole)
        cancel_button = buttons.addButton("Cancel", QDialogButtonBox.ButtonRole.RejectRole)
        save_button.setMinimumHeight(40)
        cancel_button.setMinimumHeight(40)
        outer.addWidget(buttons)
        cancel_button.clicked.connect(dialog.reject)

        def save():
            if not edits["location"].text().strip():
                QMessageBox.warning(dialog, "Required", "Location is required.")
                return
            if not observation.toPlainText().strip():
                QMessageBox.warning(dialog, "Required", "Observation description is required.")
                return
            try:
                number = next_number("HSE-OBS", "observations")
                db.execute("""
                    INSERT INTO observations (
                        number, obs_date, obs_time, project, company, location, area,
                        responsible, designation, employee_id, observer, observer_designation,
                        observer_id, observer_company, obs_type, category, subcategory,
                        observation, immediate_action, corrective_action, preventive_action,
                        priority, target_date, status, created_at
                    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """, (
                    number, today(), now_time(), edits["project"].text(), edits["company"].text(),
                    edits["location"].text(), edits["area"].text(), edits["responsible"].text(),
                    edits["designation"].text(), edits["employee_id"].text(), edits["observer"].text(),
                    edits["observer_designation"].text(), edits["observer_id"].text(),
                    edits["observer_company"].text(), obs_type.currentText(), category.currentText(),
                    edits["subcategory"].text(), observation.toPlainText(), immediate.toPlainText(),
                    corrective.toPlainText(), preventive.toPlainText(), priority.currentText(),
                    target.date().toString("yyyy-MM-dd"), status.currentText(), datetime.now().isoformat()
                ))
                row = db.fetchone("SELECT id FROM observations WHERE number=?", (number,))
                if row and attachment_paths:
                    copy_attachments(attachment_paths, number, "observation_attachments", row["id"])
                dialog.accept()
                refresh()
            except Exception as e:
                logging.exception("Observation save failed")
                QMessageBox.critical(dialog, "Save Error", f"Unable to save observation.\n\n{e}")

        save_button.clicked.connect(save)
        dialog.exec()

    def incidents(self):
        w, layout = self.page("Accident / Incident Investigation")
        banner = QHBoxLayout()
        logo = QLabel(); logo.setFixedSize(80, 60); logo.setAlignment(Qt.AlignmentFlag.AlignCenter)
        logo_path = db.setting("company_logo", "")
        if logo_path and Path(logo_path).exists():
            logo.setPixmap(QPixmap(logo_path).scaled(70, 55, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))
        banner.addWidget(logo)
        banner.addWidget(QLabel(f"<b>{safe(company_name())}</b><br>Accident / Incident Investigation Register"), 1)
        layout.addLayout(banner)
        toolbar = QHBoxLayout()
        add_button = QPushButton("+ New Incident")
        report_button = QPushButton("Generate Word Report")
        toolbar.addWidget(add_button); toolbar.addWidget(report_button)
        layout.addLayout(toolbar)

        table = QTableWidget()
        headers = ["ID","Number","Date","Type","Location","Project","Method","Status","Attachments"]
        table.setColumnCount(len(headers))
        table.setHorizontalHeaderLabels(headers)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        table.horizontalHeader().setStretchLastSection(True)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        layout.addWidget(table)

        def load():
            rows = db.fetchall("SELECT * FROM incidents ORDER BY id DESC")
            table.setRowCount(len(rows))
            for r, row in enumerate(rows):
                count = db.fetchone("SELECT COUNT(*) c FROM incident_attachments WHERE incident_id=?", (row["id"],))["c"]
                values = [row["id"], row["number"], row["incident_date"], row["incident_type"],
                          row["location"], row["project"], row["investigation_method"],
                          row["status"], count]
                for c, value in enumerate(values):
                    table.setItem(r, c, QTableWidgetItem(safe(value)))

        def generate_selected():
            r = table.currentRow()
            if r < 0:
                QMessageBox.warning(self, "Word Report", "Select an investigation first."); return
            incident_id = int(table.item(r,0).text())
            number = table.item(r,1).text()
            path, _ = QFileDialog.getSaveFileName(
                self, "Save Investigation Report", f"{number}_Investigation_Report.docx",
                "Word Document (*.docx)"
            )
            if not path: return
            try:
                generate_incident_docx(incident_id, path)
                QMessageBox.information(self, "Report Created", "Professional Word investigation report created successfully.")
            except Exception as e:
                logging.exception("Incident Word report failed")
                QMessageBox.critical(self, "Report Error", str(e))

        add_button.clicked.connect(lambda: self.incident_form(load))
        report_button.clicked.connect(generate_selected)
        load()



    def incident_form(self, refresh):
        dialog = QDialog(self)
        dialog.setWindowTitle("New Accident / Incident Investigation")
        dialog.resize(1000, 900)
        dialog.setMinimumSize(920, 760)
        outer = QVBoxLayout(dialog)

        title = QLabel("ACCIDENT / INCIDENT INVESTIGATION")
        title.setStyleSheet("font-size:20px;font-weight:bold;color:#17365D;padding:6px;")
        outer.addWidget(title)

        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        content = QWidget()
        form = QFormLayout(content)
        form.setFieldGrowthPolicy(QFormLayout.FieldGrowthPolicy.ExpandingFieldsGrow)
        edits = {}

        def add_edit(name):
            e = QLineEdit(); edits[name] = e
            form.addRow(name.replace("_", " ").title() + ":", e)
            return e

        for name in ["incident_time", "location", "project", "company", "department", "activity",
                     "person_involved", "employee_id", "designation", "supervisor", "witnesses", "equipment"]:
            add_edit(name)

        incident_type = QComboBox(); incident_type.addItems(INCIDENT_TYPES)
        form.addRow("Type of Incident:", incident_type)

        description = QTextEdit(); description.setMinimumHeight(110)
        form.addRow("Incident Description:", description)
        immediate = QTextEdit(); immediate.setMinimumHeight(80)
        form.addRow("Immediate Action:", immediate)
        consequences = QTextEdit(); consequences.setMinimumHeight(80)
        form.addRow("Actual Consequences:", consequences)
        potential = QTextEdit(); potential.setMinimumHeight(80)
        form.addRow("Potential Consequences:", potential)

        method = QComboBox(); method.addItems(INVESTIGATION_METHODS)
        form.addRow("Investigation Method:", method)

        whys = []
        for i in range(1, 6):
            e = QLineEdit(); whys.append(e)
            form.addRow(f"Why {i}:", e)

        direct = QTextEdit(); direct.setMinimumHeight(80)
        form.addRow("Direct Cause:", direct)
        contributing = QTextEdit(); contributing.setMinimumHeight(80)
        form.addRow("Contributing Factors:", contributing)
        root = QTextEdit(); root.setMinimumHeight(80)
        form.addRow("Root Cause:", root)
        corrective = QTextEdit(); corrective.setMinimumHeight(80)
        form.addRow("Corrective Action:", corrective)
        preventive = QTextEdit(); preventive.setMinimumHeight(80)
        form.addRow("Preventive Action:", preventive)

        attachment_paths = []
        attachment_label = QLabel("No evidence files selected.")
        attachment_label.setWordWrap(True)
        attach_button = QPushButton("Attach Evidence Files")
        attach_button.setMinimumHeight(36)

        def choose_files():
            paths, _ = QFileDialog.getOpenFileNames(
                dialog, "Select Investigation Evidence", "",
                "Evidence Files (*.png *.jpg *.jpeg *.bmp *.pdf *.doc *.docx *.xls *.xlsx *.txt *.mp4 *.avi *.mov);;All Files (*)"
            )
            if paths:
                attachment_paths.clear(); attachment_paths.extend(paths)
                attachment_label.setText("\n".join(Path(p).name for p in paths))

        attach_button.clicked.connect(choose_files)
        form.addRow("Evidence / Attachments:", attach_button)
        form.addRow("Selected Files:", attachment_label)

        scroll_area.setWidget(content)
        outer.addWidget(scroll_area, 1)

        buttons = QDialogButtonBox()
        save_button = buttons.addButton("Save Investigation", QDialogButtonBox.ButtonRole.AcceptRole)
        cancel_button = buttons.addButton("Cancel", QDialogButtonBox.ButtonRole.RejectRole)
        save_button.setMinimumHeight(40); cancel_button.setMinimumHeight(40)
        outer.addWidget(buttons)
        cancel_button.clicked.connect(dialog.reject)

        def save():
            if not edits["location"].text().strip():
                QMessageBox.warning(dialog, "Required", "Location is required.")
                return
            try:
                number = next_number("HSE-INC", "incidents")
                db.execute("""
                    INSERT INTO incidents (
                        number, incident_date, incident_time, location, project, company, department,
                        activity, incident_type, person_involved, employee_id, designation, supervisor,
                        witnesses, description, immediate_action, consequences, potential_consequences,
                        equipment, investigation_method, why1, why2, why3, why4, why5,
                        direct_cause, contributing_factors, root_cause, corrective_action,
                        preventive_action, created_at
                    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """, (
                    number, today(), edits["incident_time"].text(), edits["location"].text(),
                    edits["project"].text(), edits["company"].text(), edits["department"].text(),
                    edits["activity"].text(), incident_type.currentText(), edits["person_involved"].text(),
                    edits["employee_id"].text(), edits["designation"].text(), edits["supervisor"].text(),
                    edits["witnesses"].text(), description.toPlainText(), immediate.toPlainText(),
                    consequences.toPlainText(), potential.toPlainText(), edits["equipment"].text(),
                    method.currentText(), *[x.text() for x in whys], direct.toPlainText(),
                    contributing.toPlainText(), root.toPlainText(), corrective.toPlainText(),
                    preventive.toPlainText(), datetime.now().isoformat()
                ))
                row = db.fetchone("SELECT id FROM incidents WHERE number=?", (number,))
                if row and attachment_paths:
                    copy_attachments(attachment_paths, number, "incident_attachments", row["id"])
                dialog.accept()
                refresh()
            except Exception as e:
                logging.exception("Incident save failed")
                QMessageBox.critical(dialog, "Save Error", f"Unable to save investigation.\n\n{e}")

        save_button.clicked.connect(save)
        dialog.exec()

    def audits(self):
        w, layout = self.page("Audit Register")
        toolbar=QHBoxLayout()
        add_button=QPushButton("+ New Audit")
        report_button=QPushButton("Export Word")
        toolbar.addWidget(add_button); toolbar.addWidget(report_button)
        layout.addLayout(toolbar)

        table=QTableWidget()
        headers=["ID","Number","Date","Standard","Audit Type","Project","Auditor","Status","Attachments"]
        table.setColumnCount(len(headers)); table.setHorizontalHeaderLabels(headers)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        table.horizontalHeader().setStretchLastSection(True)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        layout.addWidget(table)

        def load():
            rows=db.fetchall("SELECT * FROM audits ORDER BY id DESC")
            table.setRowCount(len(rows))
            for r,row in enumerate(rows):
                count=db.fetchone("SELECT COUNT(*) c FROM audit_attachments WHERE audit_id=?",(row["id"],))["c"]
                vals=[row["id"],row["number"],row["audit_date"],row["standard"],row["audit_type"],
                      row["project"],row["lead_auditor"],row["status"],count]
                for c,v in enumerate(vals): table.setItem(r,c,QTableWidgetItem(safe(v)))

        def export_selected():
            r=table.currentRow()
            if r<0: QMessageBox.warning(self,"Word Export","Select an audit first."); return
            table_name="audits"
            path,_=QFileDialog.getSaveFileName(self,"Save Audit Word Report","audit_report.docx","Word Document (*.docx)")
            if not path:return
            try:
                generate_table_docx(table_name,path)
                QMessageBox.information(self,"Export Complete","Word report created successfully.")
            except Exception as e:
                QMessageBox.critical(self,"Export Error",str(e))
        add_button.clicked.connect(lambda:self.audit_form(load))
        report_button.clicked.connect(export_selected)
        load()



    def audit_form(self, refresh):
        dialog = QDialog(self)
        dialog.setWindowTitle("New Audit")
        dialog.resize(1000, 900)
        dialog.setMinimumSize(920, 760)
        outer = QVBoxLayout(dialog)

        title = QLabel("AUDIT REGISTER - NEW AUDIT")
        title.setStyleSheet("font-size:20px;font-weight:bold;color:#17365D;padding:6px;")
        outer.addWidget(title)

        scroll_area = QScrollArea(); scroll_area.setWidgetResizable(True)
        content = QWidget(); form = QFormLayout(content)
        form.setFieldGrowthPolicy(QFormLayout.FieldGrowthPolicy.ExpandingFieldsGrow)
        edits = {}

        def edit(name):
            e = QLineEdit(); edits[name] = e
            form.addRow(name.replace("_", " ").title() + ":", e)
            return e

        standard = QComboBox(); standard.addItems(["ISO 45001", "ISO 14001"])
        form.addRow("Standard:", standard)
        audit_type = QComboBox(); audit_type.addItems(AUDIT_TYPES)
        form.addRow("Audit Type:", audit_type)
        for name in ["project", "location", "department", "auditor", "lead_auditor", "auditee", "start_time", "end_time"]:
            edit(name)

        scope = QTextEdit(); scope.setMinimumHeight(70); form.addRow("Scope:", scope)
        objective = QTextEdit(); objective.setMinimumHeight(70); form.addRow("Objective:", objective)
        criteria = QTextEdit(); criteria.setMinimumHeight(70); form.addRow("Audit Criteria:", criteria)

        clause = QComboBox()
        sub_clause = QComboBox()
        form.addRow("ISO Clause:", clause)
        form.addRow("ISO Sub-Clause:", sub_clause)

        def load_clauses():
            clause.blockSignals(True)
            clause.clear()
            clause.addItems(list(ISO_CLAUSES.get(standard.currentText(), {}).keys()))
            clause.blockSignals(False)
            load_subclauses()

        def load_subclauses():
            sub_clause.clear()
            items = ISO_CLAUSES.get(standard.currentText(), {}).get(clause.currentText(), [])
            sub_clause.addItems(items)

        standard.currentTextChanged.connect(load_clauses)
        clause.currentTextChanged.connect(load_subclauses)
        load_clauses()

        requirement = QTextEdit(); requirement.setMinimumHeight(70); form.addRow("Requirement / Criterion:", requirement)
        finding_type = QComboBox(); finding_type.addItems(FINDING_TYPES); form.addRow("Finding Type:", finding_type)
        finding = QTextEdit(); finding.setMinimumHeight(90); form.addRow("Finding / Evidence:", finding)
        risk = QTextEdit(); risk.setMinimumHeight(70); form.addRow("Risk / Impact:", risk)
        corrective = QTextEdit(); corrective.setMinimumHeight(70); form.addRow("Corrective Action:", corrective)
        responsible = QLineEdit(); form.addRow("Responsible Person:", responsible)
        target = QLineEdit(); form.addRow("Target Date:", target)

        attachment_paths = []; attachment_label = QLabel("No evidence files selected."); attachment_label.setWordWrap(True)
        attach = QPushButton("Attach Audit Evidence Files"); attach.setMinimumHeight(36)
        def choose_files():
            paths, _ = QFileDialog.getOpenFileNames(
                dialog, "Select Audit Evidence", "",
                "Evidence Files (*.png *.jpg *.jpeg *.bmp *.pdf *.doc *.docx *.xls *.xlsx *.txt *.mp4 *.avi *.mov);;All Files (*)"
            )
            if paths:
                attachment_paths.clear(); attachment_paths.extend(paths)
                attachment_label.setText("\n".join(Path(p).name for p in paths))
        attach.clicked.connect(choose_files)
        form.addRow("Evidence / Attachments:", attach); form.addRow("Selected Files:", attachment_label)

        scroll_area.setWidget(content); outer.addWidget(scroll_area, 1)
        buttons = QDialogButtonBox()
        save_button = buttons.addButton("Save Audit", QDialogButtonBox.ButtonRole.AcceptRole)
        cancel_button = buttons.addButton("Cancel", QDialogButtonBox.ButtonRole.RejectRole)
        save_button.setMinimumHeight(40); cancel_button.setMinimumHeight(40)
        outer.addWidget(buttons); cancel_button.clicked.connect(dialog.reject)

        def save():
            try:
                number = next_number("HSE-AUD", "audits")
                db.execute("""
                    INSERT INTO audits (number,audit_date,audit_type,standard,project,location,department,
                        auditor,lead_auditor,auditee,scope,objective,criteria,start_time,end_time,created_at)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """, (
                    number, today(), audit_type.currentText(), standard.currentText(), edits["project"].text(),
                    edits["location"].text(), edits["department"].text(), edits["auditor"].text(),
                    edits["lead_auditor"].text(), edits["auditee"].text(), scope.toPlainText(),
                    objective.toPlainText(), criteria.toPlainText(), edits["start_time"].text(),
                    edits["end_time"].text(), datetime.now().isoformat()
                ))
                audit = db.fetchone("SELECT id FROM audits WHERE number=?", (number,))
                db.execute("""
                    INSERT INTO audit_findings (audit_id,clause,sub_clause,requirement,finding_type,
                        observation,evidence,risk_impact,corrective_action,responsible,target_date)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?)
                """, (
                    audit["id"], clause.currentText(), sub_clause.currentText(), requirement.toPlainText(),
                    finding_type.currentText(), finding.toPlainText(), finding.toPlainText(), risk.toPlainText(),
                    corrective.toPlainText(), responsible.text(), target.text()
                ))
                if attachment_paths:
                    copy_attachments(attachment_paths, number, "audit_attachments", audit["id"])
                dialog.accept(); refresh()
            except Exception as e:
                logging.exception("Audit save failed")
                QMessageBox.critical(dialog, "Save Error", f"Unable to save audit.\n\n{e}")

        save_button.clicked.connect(save)
        dialog.exec()

    def capa(self):
        w,layout=self.page("CAPA Register")
        toolbar=QHBoxLayout(); add_button=QPushButton("+ New CAPA"); word_button=QPushButton("Export Word")
        toolbar.addWidget(add_button); toolbar.addWidget(word_button); layout.addLayout(toolbar)
        table=QTableWidget()
        headers=["ID","Number","Source","Reference","Priority","Responsible","Target Date","Status","Overdue","Attachments"]
        table.setColumnCount(len(headers)); table.setHorizontalHeaderLabels(headers)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        table.horizontalHeader().setStretchLastSection(True); table.setSelectionBehavior(QAbstractItemView.SelectRows)
        layout.addWidget(table)
        def load():
            rows=db.fetchall("SELECT * FROM capa ORDER BY id DESC"); table.setRowCount(len(rows))
            for r,row in enumerate(rows):
                count=db.fetchone("SELECT COUNT(*) c FROM capa_attachments WHERE capa_id=?",(row["id"],))["c"]
                vals=[row["id"],row["number"],row["source"],row["reference_number"],row["priority"],row["responsible"],
                      row["target_date"],row["status"],"OVERDUE" if overdue(row["target_date"],row["status"]) else "",count]
                for c,v in enumerate(vals): table.setItem(r,c,QTableWidgetItem(safe(v)))
        def export_word():
            path,_=QFileDialog.getSaveFileName(self,"Save CAPA Word Report","capa_report.docx","Word Document (*.docx)")
            if not path:return
            try: generate_table_docx("capa",path); QMessageBox.information(self,"Export Complete","Word report created successfully.")
            except Exception as e: QMessageBox.critical(self,"Export Error",str(e))
        add_button.clicked.connect(lambda:self.capa_form(load)); word_button.clicked.connect(export_word); load()



    def capa_form(self, refresh):
        dialog = QDialog(self)
        dialog.setWindowTitle("New CAPA")
        dialog.resize(950, 850)
        dialog.setMinimumSize(900, 720)
        outer = QVBoxLayout(dialog)

        title = QLabel("CORRECTIVE AND PREVENTIVE ACTION (CAPA)")
        title.setStyleSheet("font-size:20px;font-weight:bold;color:#17365D;padding:6px;")
        outer.addWidget(title)

        scroll_area = QScrollArea(); scroll_area.setWidgetResizable(True)
        content = QWidget(); form = QFormLayout(content)
        form.setFieldGrowthPolicy(QFormLayout.FieldGrowthPolicy.ExpandingFieldsGrow)

        source = QComboBox(); source.addItems(CAPA_SOURCES); form.addRow("Source:", source)
        reference = QLineEdit(); form.addRow("Reference Number:", reference)
        finding = QTextEdit(); finding.setMinimumHeight(90); form.addRow("Finding:", finding)
        root = QTextEdit(); root.setMinimumHeight(90); form.addRow("Root Cause:", root)
        corrective = QTextEdit(); corrective.setMinimumHeight(90); form.addRow("Corrective Action:", corrective)
        preventive = QTextEdit(); preventive.setMinimumHeight(90); form.addRow("Preventive Action:", preventive)
        responsible = QLineEdit(); form.addRow("Responsible Person:", responsible)
        priority = QComboBox(); priority.addItems(PRIORITIES); form.addRow("Priority:", priority)
        target = QLineEdit(); form.addRow("Target Date:", target)
        status = QComboBox(); status.addItems(STATUSES); form.addRow("Status:", status)
        verification = QTextEdit(); verification.setMinimumHeight(70); form.addRow("Verification:", verification)
        evidence = QTextEdit(); evidence.setMinimumHeight(70); form.addRow("Closeout Evidence:", evidence)

        attachment_paths = []; attachment_label = QLabel("No evidence files selected."); attachment_label.setWordWrap(True)
        attach = QPushButton("Attach CAPA Evidence Files"); attach.setMinimumHeight(36)
        def choose_files():
            paths, _ = QFileDialog.getOpenFileNames(
                dialog, "Select CAPA Evidence", "",
                "Evidence Files (*.png *.jpg *.jpeg *.bmp *.pdf *.doc *.docx *.xls *.xlsx *.txt *.mp4 *.avi *.mov);;All Files (*)"
            )
            if paths:
                attachment_paths.clear(); attachment_paths.extend(paths)
                attachment_label.setText("\n".join(Path(p).name for p in paths))
        attach.clicked.connect(choose_files)
        form.addRow("Evidence / Attachments:", attach); form.addRow("Selected Files:", attachment_label)

        scroll_area.setWidget(content); outer.addWidget(scroll_area, 1)
        buttons = QDialogButtonBox()
        save_button = buttons.addButton("Save CAPA", QDialogButtonBox.ButtonRole.AcceptRole)
        cancel_button = buttons.addButton("Cancel", QDialogButtonBox.ButtonRole.RejectRole)
        save_button.setMinimumHeight(40); cancel_button.setMinimumHeight(40)
        outer.addWidget(buttons); cancel_button.clicked.connect(dialog.reject)

        def save():
            try:
                number = next_number("HSE-CAPA", "capa")
                db.execute("""
                    INSERT INTO capa (number,source,reference_number,finding,root_cause,corrective_action,
                        preventive_action,responsible,priority,target_date,status,verification,closeout_evidence,created_at)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """, (
                    number, source.currentText(), reference.text(), finding.toPlainText(), root.toPlainText(),
                    corrective.toPlainText(), preventive.toPlainText(), responsible.text(), priority.currentText(),
                    target.text(), status.currentText(), verification.toPlainText(), evidence.toPlainText(),
                    datetime.now().isoformat()
                ))
                row = db.fetchone("SELECT id FROM capa WHERE number=?", (number,))
                if row and attachment_paths:
                    copy_attachments(attachment_paths, number, "capa_attachments", row["id"])
                dialog.accept(); refresh()
            except Exception as e:
                logging.exception("CAPA save failed")
                QMessageBox.critical(dialog, "Save Error", f"Unable to save CAPA.\n\n{e}")

        save_button.clicked.connect(save)
        dialog.exec()

    def reports(self):
        w, layout = self.page("Reports & Export")
        banner = QHBoxLayout()
        logo_path = db.setting("company_logo", "")
        logo = QLabel()
        logo.setFixedSize(90, 70)
        logo.setAlignment(Qt.AlignmentFlag.AlignCenter)
        if logo_path and Path(logo_path).exists():
            pix = QPixmap(logo_path)
            logo.setPixmap(pix.scaled(80, 65, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))
        banner.addWidget(logo)
        company_label = QLabel(f"<b>{safe(company_name())}</b><br>Reports & Export<br>Project: {safe(db.setting('project_name',''))}")
        company_label.setStyleSheet("font-size:16px;padding:6px;")
        banner.addWidget(company_label, 1)
        open_folder = QPushButton("Open Reports Folder")
        open_folder.clicked.connect(self.open_reports_folder)
        banner.addWidget(open_folder)
        layout.addLayout(banner)

        info = QLabel("Select a register and format. The report is extracted from the saved HSE database and can be saved to any folder on the computer.")
        info.setWordWrap(True)
        layout.addWidget(info)

        for title, table_name in [
            ("HSE Inspection / Observation Report", "observations"),
            ("Incident Investigation Register Report", "incidents"),
            ("Audit Register Report", "audits"),
            ("CAPA Register Report", "capa")
        ]:
            box = QGroupBox(title)
            row = QHBoxLayout(box)
            for label, handler in [
                ("Download CSV", lambda t=table_name: self.export_csv(t)),
                ("Download Excel", lambda t=table_name: self.export_excel(t)),
                ("Download PDF", lambda t=table_name: self.export_pdf(t)),
                ("Download Word", lambda t=table_name: self.export_docx(t))
            ]:
                b = QPushButton(label)
                b.setMinimumHeight(38)
                row.addWidget(b)
                b.clicked.connect(handler)
            layout.addWidget(box)

        incident_box = QGroupBox("Professional Incident Investigation Word Report")
        ir = QHBoxLayout(incident_box)
        incident_combo = QComboBox()
        incidents = db.fetchall("SELECT id,number FROM incidents ORDER BY id DESC")
        for row in incidents:
            incident_combo.addItem(safe(row["number"]), row["id"])
        ir.addWidget(QLabel("Investigation:"))
        ir.addWidget(incident_combo, 1)
        b = QPushButton("Download Investigation Report")
        b.setMinimumHeight(38)
        ir.addWidget(b)

        def generate():
            if incident_combo.count() == 0:
                QMessageBox.information(self, "No Incidents", "No incident investigations are available.")
                return
            incident_id = incident_combo.currentData()
            number = incident_combo.currentText()
            default_path = str(REPORT_DIR / f"{number}_Investigation_Report.docx")
            path, _ = QFileDialog.getSaveFileName(
                self, "Save Investigation Report", default_path, "Word Document (*.docx)"
            )
            if not path:
                return
            try:
                generate_incident_docx(int(incident_id), path)
                self.show_export_success(path)
            except Exception as e:
                logging.exception("Incident report failed")
                QMessageBox.critical(self, "Report Error", str(e))

        b.clicked.connect(generate)
        layout.addWidget(incident_box)
        layout.addStretch(1)

    def show_export_success(self, path):
        msg = QMessageBox(self)
        msg.setWindowTitle("Export Complete")
        msg.setIcon(QMessageBox.Icon.Information)
        msg.setText("Report created successfully.")
        msg.setInformativeText(str(path))
        open_button = msg.addButton("Open Folder", QMessageBox.ButtonRole.ActionRole)
        msg.addButton("OK", QMessageBox.ButtonRole.AcceptRole)
        msg.exec()
        if msg.clickedButton() == open_button:
            self.open_path(Path(path).parent)

    def open_path(self, path):
        try:
            path = Path(path)
            if sys.platform.startswith("win"):
                os.startfile(str(path))
            elif sys.platform == "darwin":
                os.system(f'open "{path}"')
            else:
                os.system(f'xdg-open "{path}"')
        except Exception as e:
            logging.exception("Unable to open path")
            QMessageBox.warning(self, "Open Folder", f"Unable to open folder.\n\n{e}")

    def open_reports_folder(self):
        REPORT_DIR.mkdir(parents=True, exist_ok=True)
        self.open_path(REPORT_DIR)

    def _export_path(self, filename, title, file_filter):
        REPORT_DIR.mkdir(parents=True, exist_ok=True)
        return QFileDialog.getSaveFileName(self, title, str(REPORT_DIR / filename), file_filter)[0]

    @staticmethod
    def table_data_static(table_name):
        allowed = {"observations", "incidents", "audits", "capa"}
        if table_name not in allowed:
            raise ValueError("Invalid report table.")
        rows = db.fetchall(f"SELECT * FROM {table_name} ORDER BY id DESC")
        if not rows:
            return [], []
        columns = list(rows[0].keys())
        return columns, [[safe(row[c]) for c in columns] for row in rows]

    def table_data(self, table_name):
        return self.table_data_static(table_name)

    def export_csv(self, table_name):
        try:
            columns, data = self.table_data(table_name)
            if not columns:
                QMessageBox.information(self, "No Data", "There is no data to export.")
                return
            path = self._export_path(f"{table_name}_report.csv", "Download CSV Report", "CSV Files (*.csv)")
            if not path:
                return
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f); writer.writerow(columns); writer.writerows(data)
            self.show_export_success(path)
        except Exception as e:
            logging.exception("CSV export failed")
            QMessageBox.critical(self, "Export Error", str(e))

    def export_excel(self, table_name):
        try:
            columns, data = self.table_data(table_name)
            if not columns:
                QMessageBox.information(self, "No Data", "There is no data to export.")
                return
            path = self._export_path(f"{table_name}_report.xlsx", "Download Excel Report", "Excel Files (*.xlsx)")
            if not path:
                return
            wb = Workbook(); ws = wb.active; ws.title = table_name[:31]; ws.append(columns)
            for row in data: ws.append(row)
            for cell in ws[1]:
                font = copy(cell.font); font.bold = True; cell.font = font
            ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
            wb.save(path)
            self.show_export_success(path)
        except Exception as e:
            logging.exception("Excel export failed")
            QMessageBox.critical(self, "Export Error", str(e))

    def export_pdf(self, table_name):
        try:
            columns, data = self.table_data(table_name)
            if not columns:
                QMessageBox.information(self, "No Data", "There is no data to export.")
                return
            path = self._export_path(f"{table_name}_report.pdf", "Download PDF Report", "PDF Files (*.pdf)")
            if not path:
                return
            styles = getSampleStyleSheet()
            doc = SimpleDocTemplate(path, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=25, bottomMargin=25)
            story = []
            logo_path = db.setting("company_logo", "")
            if logo_path and Path(logo_path).exists():
                try:
                    from reportlab.platypus import Image
                    story.append(Image(logo_path, width=70, height=70))
                except Exception:
                    pass
            story.append(Paragraph(company_name(), styles["Title"]))
            project = db.setting("project_name", "")
            if project: story.append(Paragraph(project, styles["Heading2"]))
            prefix = db.setting("document_prefix", "HSE")
            story.append(Paragraph(f"{table_name.title()} Report | Document Prefix: {prefix}", styles["Heading2"]))
            story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]))
            story.append(Spacer(1, 12))
            max_columns = min(len(columns), 10)
            pdf_data = [columns[:max_columns]] + [r[:max_columns] for r in data[:500]]
            t = Table(pdf_data, repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#17365D")),
                ("TEXTCOLOR", (0,0), (-1,0), colors.white),
                ("GRID", (0,0), (-1,-1), 0.4, colors.grey),
                ("FONTSIZE", (0,0), (-1,-1), 6),
                ("VALIGN", (0,0), (-1,-1), "TOP")
            ]))
            story.append(t)
            footer = db.setting("report_footer", "")
            if footer: story.extend([Spacer(1,10), Paragraph(footer, styles["Normal"])])
            doc.build(story)
            self.show_export_success(path)
        except Exception as e:
            logging.exception("PDF export failed")
            QMessageBox.critical(self, "Export Error", str(e))

    def export_docx(self, table_name):
        try:
            columns, data = self.table_data(table_name)
            if not columns:
                QMessageBox.information(self, "No Data", "There is no data to export.")
                return
            path = self._export_path(f"{table_name}_report.docx", "Download Word Report", "Word Documents (*.docx)")
            if not path:
                return
            generate_table_docx(table_name, path)
            self.show_export_success(path)
        except Exception as e:
            logging.exception("Word report failed")
            QMessageBox.critical(self, "Export Error", str(e))

    def export_table(
        self,
        table,
        filename
    ):

        path, _ = QFileDialog.getSaveFileName(
            self,
            "Export CSV",
            f"{filename}.csv",
            "CSV (*.csv)"
        )

        if not path:
            return

        with open(
            path,
            "w",
            newline="",
            encoding="utf-8-sig"
        ) as f:

            writer = csv.writer(f)

            headers = []

            for c in range(
                table.columnCount()
            ):
                headers.append(
                    table.horizontalHeaderItem(
                        c
                    ).text()
                )

            writer.writerow(headers)

            for r in range(
                table.rowCount()
            ):

                row = []

                for c in range(
                    table.columnCount()
                ):

                    item = table.item(r, c)

                    row.append(
                        item.text()
                        if item
                        else ""
                    )

                writer.writerow(row)

        QMessageBox.information(
            self,
            "Export Complete",
            "CSV exported successfully."
        )

    # ========================================================
    # MASTER DATA
    # ========================================================


    def master_data(self):
        w,layout=self.page("Master Data")
        layout.addWidget(QLabel("Manage employees, companies, projects, categories and company profile."))
        profile=QPushButton("Company Profile / Logo / Document Number")
        profile.clicked.connect(lambda:self.settings_page())
        layout.addWidget(profile)
        for title,table_name in [
            ("Employees","employees"),("Companies","companies"),
            ("Projects / Locations","projects"),("Observation Categories","categories")
        ]:
            button=QPushButton(title)
            button.clicked.connect(lambda checked=False,t=table_name,n=title:self.master_editor(t,n))
            layout.addWidget(button)


    def master_editor(
        self,
        table_name,
        title
    ):

        dialog = QDialog(self)
        dialog.setWindowTitle(title)
        dialog.resize(650, 550)

        layout = QVBoxLayout(dialog)

        table = QTableWidget()

        layout.addWidget(table)

        form = QFormLayout()

        if table_name == "employees":

            fields = [
                ("Employee ID", "employee_id"),
                ("Name", "name"),
                ("Designation", "designation"),
                ("Company", "company"),
                ("Department", "department"),
                ("Contact", "contact")
            ]

        elif table_name == "companies":

            fields = [
                ("Company Name", "name"),
                ("Company Type", "company_type"),
                ("Contact Person", "contact_person"),
                ("Contact Number", "contact_number")
            ]

        elif table_name == "projects":

            fields = [
                ("Project", "project"),
                ("Area", "area"),
                ("Location", "location"),
                ("Description", "description")
            ]

        else:

            fields = [
                ("Category", "category"),
                ("Subcategory", "subcategory")
            ]

        inputs = {}

        for label, key in fields:

            e = QLineEdit()

            inputs[key] = e

            form.addRow(
                label,
                e
            )

        layout.addLayout(form)

        add = QPushButton("Add")
        layout.addWidget(add)

        close = QPushButton("Close")
        layout.addWidget(close)

        close.clicked.connect(
            dialog.accept
        )

        def load():

            rows = db.fetchall(
                f"SELECT * FROM {table_name}"
            )

            if not rows:
                table.setRowCount(0)
                return

            columns = list(rows[0].keys())

            table.setColumnCount(
                len(columns)
            )

            table.setHorizontalHeaderLabels(
                columns
            )

            table.setRowCount(
                len(rows)
            )

            for r, row in enumerate(rows):

                for c, col in enumerate(columns):

                    table.setItem(
                        r,
                        c,
                        QTableWidgetItem(
                            safe(row[col])
                        )
                    )

            table.horizontalHeader().setSectionResizeMode(
                QHeaderView.Stretch
            )

        def save():

            values = [
                inputs[key].text()
                for _, key in fields
            ]

            if not values[0].strip():
                return

            placeholders = ",".join(
                ["?"] * len(values)
            )

            columns = ",".join(
                [key for _, key in fields]
            )

            db.execute(
                f"""
                INSERT INTO {table_name}
                ({columns})
                VALUES ({placeholders})
                """,
                values
            )

            for e in inputs.values():
                e.clear()

            load()

        add.clicked.connect(save)

        load()

        dialog.exec()

    # ========================================================
    # SETTINGS
    # ========================================================


    def settings_page(self):
        w,layout=self.page("Settings")
        form=QFormLayout()
        company=QLineEdit(db.setting("company_name",""))
        project=QLineEdit(db.setting("project_name",""))
        location=QLineEdit(db.setting("default_location",""))
        observer=QLineEdit(db.setting("default_observer",""))
        footer=QLineEdit(db.setting("report_footer",""))
        prefix=QLineEdit(db.setting("document_prefix","HSE"))
        logo_path=QLineEdit(db.setting("company_logo",""))
        logo_path.setReadOnly(True)
        logo_button=QPushButton("Choose Company Logo")
        logo_preview=QLabel()
        logo_preview.setMinimumHeight(80)
        logo_preview.setAlignment(Qt.AlignmentFlag.AlignCenter)

        form.addRow("Company Name",company)
        form.addRow("Project Name",project)
        form.addRow("Default Location",location)
        form.addRow("Default Observer",observer)
        form.addRow("Document Number / Prefix",prefix)
        form.addRow("Report Footer",footer)
        logo_row=QHBoxLayout(); logo_row.addWidget(logo_path); logo_row.addWidget(logo_button)
        form.addRow("Company Logo",logo_row)
        layout.addLayout(form); layout.addWidget(logo_preview)

        def choose_logo():
            path,_=QFileDialog.getOpenFileName(self,"Select Company Logo","","Images (*.png *.jpg *.jpeg *.bmp)")
            if not path:return
            try:
                src=Path(path); dest=ATTACH_DIR / f"company_logo{src.suffix.lower()}"
            except Exception:
                dest=ATTACH_DIR / f"company_logo{Path(path).suffix.lower()}"
            try:
                for old in ATTACH_DIR.glob("company_logo.*"):
                    if old != dest: old.unlink(missing_ok=True)
                shutil.copy2(path,dest)
                logo_path.setText(str(dest))
                pix=QPixmap(str(dest))
                logo_preview.setPixmap(pix.scaled(180,80,Qt.AspectRatioMode.KeepAspectRatio,Qt.TransformationMode.SmoothTransformation))
            except Exception as e:
                QMessageBox.critical(self,"Logo Error",str(e))

        logo_button.clicked.connect(choose_logo)
        existing_logo=db.setting("company_logo","")
        if existing_logo and Path(existing_logo).exists():
            pix=QPixmap(existing_logo)
            logo_preview.setPixmap(pix.scaled(180,80,Qt.AspectRatioMode.KeepAspectRatio,Qt.TransformationMode.SmoothTransformation))

        save=QPushButton("Save Settings"); layout.addWidget(save)
        save.clicked.connect(lambda:self.save_settings(company,project,location,observer,footer,prefix,logo_path))
        backup=QPushButton("Backup Complete Application Data"); layout.addWidget(backup); backup.clicked.connect(self.backup)
        restore=QPushButton("Restore Application Backup"); layout.addWidget(restore); restore.clicked.connect(self.restore)
        about=QPushButton("About"); layout.addWidget(about)
        about.clicked.connect(lambda:QMessageBox.information(self,"About",f"{APP_NAME}\nVersion {APP_VERSION}\n\nOffline HSE Management System"))

    def save_settings(self, company, project, location, observer, footer, prefix=None, logo_path=None):
        db.set_setting("company_name",company.text())
        db.set_setting("project_name",project.text())
        db.set_setting("default_location",location.text())
        db.set_setting("default_observer",observer.text())
        db.set_setting("report_footer",footer.text())
        if prefix is not None:
            db.set_setting("document_prefix",prefix.text().strip() or "HSE")
        if logo_path is not None:
            db.set_setting("company_logo",logo_path.text())
        QMessageBox.information(self,"Saved","Settings saved successfully.")
        self.dashboard()


    def backup(self):

        destination, _ = QFileDialog.getSaveFileName(
            self,
            "Save HSE Backup",
            f"HSE_Backup_{today()}.zip",
            "ZIP (*.zip)"
        )

        if not destination:
            return

        try:

            with zipfile.ZipFile(
                destination,
                "w",
                zipfile.ZIP_DEFLATED
            ) as z:

                if DB_FILE.exists():

                    z.write(
                        DB_FILE,
                        "database/hse.db"
                    )

                for file in ATTACH_DIR.rglob("*"):

                    if file.is_file():

                        z.write(
                            file,
                            f"attachments/{file.name}"
                        )

            QMessageBox.information(
                self,
                "Backup Complete",
                "HSE backup created successfully."
            )

        except Exception as e:

            logging.exception(
                "Backup failed"
            )

            QMessageBox.critical(
                self,
                "Backup Error",
                str(e)
            )

    # --------------------------------------------------------

    def restore(self):

        source, _ = QFileDialog.getOpenFileName(
            self,
            "Select HSE Backup",
            "",
            "ZIP (*.zip)"
        )

        if not source:
            return

        answer = QMessageBox.question(
            self,
            "Restore Backup",
            "Restoring will replace the current database.\n\n"
            "Do you want to continue?"
        )

        if answer != QMessageBox.Yes:
            return

        try:

            # Close current database connection.
            db.conn.close()

            with zipfile.ZipFile(
                source,
                "r"
            ) as z:

                temp = APP_DIR / "restore_temp"

                if temp.exists():
                    shutil.rmtree(temp)

                temp.mkdir()

                z.extractall(temp)

                restored_db = (
                    temp /
                    "database" /
                    "hse.db"
                )

                if not restored_db.exists():
                    raise Exception(
                        "Backup database not found."
                    )

                shutil.copy2(
                    restored_db,
                    DB_FILE
                )

                restored_attach = (
                    temp /
                    "attachments"
                )

                if restored_attach.exists():

                    for file in restored_attach.iterdir():

                        if file.is_file():

                            shutil.copy2(
                                file,
                                ATTACH_DIR /
                                file.name
                            )

                shutil.rmtree(temp)

            QMessageBox.information(
                self,
                "Restore Complete",
                "Backup restored successfully.\n"
                "Please restart the application."
            )

        except Exception as e:

            logging.exception(
                "Restore failed"
            )

            QMessageBox.critical(
                self,
                "Restore Error",
                str(e)
            )

        finally:

            # Reopen database.
            db.conn = sqlite3.connect(
                DB_FILE
            )

            db.conn.row_factory = sqlite3.Row
            db.conn.execute(
                "PRAGMA foreign_keys = ON"
            )

    # ========================================================
    # CLOSE
    # ========================================================

    def closeEvent(self, event):

        try:
            db.conn.commit()
            db.conn.close()
        except Exception:
            pass

        event.accept()


# ============================================================
# APPLICATION START
# ============================================================

def main():

    app = QApplication(sys.argv)

    app.setApplicationName(
        APP_NAME
    )

    app.setOrganizationName(
        "HSE"
    )

    window = MainWindow()

    window.show()

    sys.exit(
        app.exec()
    )


if __name__ == "__main__":
    main()
